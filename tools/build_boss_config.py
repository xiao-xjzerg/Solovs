#!/usr/bin/env python3
"""Compile the local Boss design workbook into runtime JSON, CSV mirrors and manifest.

The workbook lives at docs/skill_config_v1.xlsx. Runtime code never parses XLSX.
This tool owns these generated files:

  assets/data/boss_config.json
  assets/data/csv/*.csv
  assets/asset-manifest.json

The hand-maintained player/global asset input is:

  assets/asset-manifest.source.json
"""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import io
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "docs" / "skill_config_v1.xlsx"
DEFAULT_MANIFEST_SOURCE = PROJECT_ROOT / "assets" / "asset-manifest.source.json"
DEFAULT_MANIFEST_OUT = PROJECT_ROOT / "assets" / "asset-manifest.json"
DEFAULT_CONFIG_OUT = PROJECT_ROOT / "assets" / "data" / "boss_config.json"
DEFAULT_CSV_DIR = PROJECT_ROOT / "assets" / "data" / "csv"

GENERATOR_VERSION = 3
CONFIG_SCHEMA_VERSION = 2
PLACEHOLDERS = {"tbd", "todo", "to be determined", "待定", "待补"}
ACTOR_REFERENCE_POINTS = {
    "actor",
    "self",
    "feet",
    "body_center",
    "hurtbox_center",
    "display_center",
}

REQUIRED_COLUMNS: dict[str, set[str]] = {
    "Skills": {
        "skill_id",
        "anim_id",
        "cast_range_min",
        "cast_range_max",
        "cooldown_ms",
        "weight",
        "damage",
        "params",
    },
    "melee_params": {
        "shape",
        "origin",
        "origin_policy",
        "origin_forward_offset",
        "facing_policy",
        "angle_deg",
        "radius",
        "origin_offset_x",
        "origin_offset_y",
        "target_reference",
    },
    "Skill_events": {"event_id", "skill_id", "trigger", "event_type", "ref_id"},
    "Skill_effects": {"effect_id", "skill_id", "trigger", "effect_type"},
    "Projectiles": {"projectile_id", "image_path", "move_mode", "target_policy"},
    "Animations": {"anim_id", "owner_id", "action", "loop", "direction_mode", "required"},
    "Animation_frames": {"anim_id", "frame", "image_path", "duration_ms"},
    "VFX_assets": {"vfx_id", "kind", "renderer", "required"},
    "SFX_assets": {"sfx_id", "path", "required"},
    "Bosses": {
        "boss_id",
        "max_hp",
        "move_speed",
        "aggro_range",
        "body_shape",
        "hurtbox_shape",
        "initial_phase",
    },
    "Boss_phases": {"boss_id", "phase", "hp_min_pct", "hp_max_pct"},
    "Phase_skills": {"boss_id", "phase", "skill_id", "enabled"},
}


class BuildFailure(Exception):
    pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build Solovs BOSS runtime configuration")
    parser.add_argument("--source", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--manifest-source", type=Path, default=DEFAULT_MANIFEST_SOURCE)
    parser.add_argument("--manifest-out", type=Path, default=DEFAULT_MANIFEST_OUT)
    parser.add_argument("--config-out", type=Path, default=DEFAULT_CONFIG_OUT)
    parser.add_argument("--csv-dir", type=Path, default=DEFAULT_CSV_DIR)
    parser.add_argument("--check", action="store_true", help="Validate without writing generated files")
    parser.add_argument("--no-csv", action="store_true", help="Skip CSV mirror generation")
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def is_placeholder(value: Any) -> bool:
    return isinstance(value, str) and value.strip().lower() in PLACEHOLDERS


def clean_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_ref(value: Any) -> str | None:
    text = clean_string(value)
    if not text:
        return None
    if len(text) >= 2 and text[0] == text[-1] and text[0] in {"'", '"'}:
        text = text[1:-1].strip()
    return text or None


def normalize_path(value: Any) -> str | None:
    text = clean_string(value)
    if not text or is_placeholder(text):
        return text
    return re.sub(r"/+", "/", text.replace("\\", "/"))


def camel_name(name: str) -> str:
    parts = re.split(r"[_\s]+", name.strip())
    if not parts:
        return name
    return parts[0].lower() + "".join(part[:1].upper() + part[1:] for part in parts[1:])


def camel_mapping(row: dict[str, Any], *, path_fields: Iterable[str] = ()) -> dict[str, Any]:
    path_fields = set(path_fields)
    result: dict[str, Any] = {}
    for key, value in row.items():
        if key in path_fields:
            value = normalize_path(value)
        elif isinstance(value, str):
            value = value.strip() or None
        result[camel_name(key)] = value
    return result


def number(value: Any, field: str, errors: list[str], *, allow_none: bool = False) -> float | int | None:
    if value is None or value == "":
        if allow_none:
            return None
        errors.append(f"{field}: required number is blank")
        return 0
    if isinstance(value, bool):
        errors.append(f"{field}: boolean is not a valid number")
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        parsed = float(str(value).strip())
        return int(parsed) if parsed.is_integer() else parsed
    except ValueError:
        if allow_none and is_placeholder(value):
            return None
        errors.append(f"{field}: expected number, got {value!r}")
        return None if allow_none else 0


def boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value or "").strip().lower() in {"true", "1", "yes", "y"}


def read_workbook(path: Path) -> tuple[Any, dict[str, list[dict[str, Any]]]]:
    if not path.exists():
        raise BuildFailure(f"Workbook does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            sheets[worksheet.title] = []
            continue
        headers = [clean_string(value) or "" for value in rows[0]]
        records: list[dict[str, Any]] = []
        for raw in rows[1:]:
            if not any(value is not None and str(value).strip() for value in raw):
                continue
            records.append({headers[index]: raw[index] for index in range(len(headers)) if headers[index]})
        sheets[worksheet.title] = records
    return workbook, sheets


def validate_columns(sheets: dict[str, list[dict[str, Any]]], errors: list[str]) -> None:
    for sheet_name, required in REQUIRED_COLUMNS.items():
        if sheet_name not in sheets:
            errors.append(f"Missing worksheet: {sheet_name}")
            continue
        rows = sheets[sheet_name]
        if not rows:
            errors.append(f"Worksheet has no data rows: {sheet_name}")
            continue
        actual = set(rows[0])
        missing = sorted(required - actual)
        if missing:
            errors.append(f"{sheet_name}: missing columns {missing}")


def unique_index(
    rows: list[dict[str, Any]],
    key: str,
    sheet_name: str,
    errors: list[str],
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(rows, start=2):
        identifier = clean_string(row.get(key))
        if not identifier:
            errors.append(f"{sheet_name}!{index}: blank {key}")
            continue
        if identifier in result:
            errors.append(f"{sheet_name}!{index}: duplicate {key}={identifier}")
            continue
        result[identifier] = row
    return result


def workbook_sheet_parameters(
    sheets: dict[str, list[dict[str, Any]]],
    reference: str | None,
    where: str,
    errors: list[str],
) -> dict[str, Any]:
    if not reference:
        return {}
    rows = sheets.get(reference)
    if rows is None:
        errors.append(f"{where}: params worksheet does not exist: {reference}")
        return {}
    if len(rows) != 1:
        errors.append(f"{where}: params worksheet {reference} must contain exactly one data row")
        return {}
    return camel_mapping(rows[0])


def infer_skill_handler(skill_id: str, params: dict[str, Any], events: list[dict[str, Any]]) -> str:
    event_types = {event.get("eventType") for event in events}
    if "projectile_spawn" in event_types:
        return "projectile"
    if "wave_start" in {event.get("trigger") for event in events}:
        return "spawner"
    if "endLocation" in params:
        return "charge"
    shape = clean_string(params.get("shape"))
    if shape in {"arc", "circle"}:
        return shape
    if skill_id == "crystal_pulse":
        return "spawner"
    if skill_id == "bull_rush":
        return "charge"
    if not params:
        return "sequence"
    return shape or "unsupported"


def compile_config(
    sheets: dict[str, list[dict[str, Any]]],
    source_path: Path,
    errors: list[str],
    warnings: list[str],
) -> dict[str, Any]:
    skill_rows = unique_index(sheets.get("Skills", []), "skill_id", "Skills", errors)
    event_rows = unique_index(sheets.get("Skill_events", []), "event_id", "Skill_events", errors)
    effect_rows = unique_index(sheets.get("Skill_effects", []), "effect_id", "Skill_effects", errors)
    projectile_rows = unique_index(sheets.get("Projectiles", []), "projectile_id", "Projectiles", errors)
    animation_rows = unique_index(sheets.get("Animations", []), "anim_id", "Animations", errors)
    vfx_rows = unique_index(sheets.get("VFX_assets", []), "vfx_id", "VFX_assets", errors)
    sfx_rows = unique_index(sheets.get("SFX_assets", []), "sfx_id", "SFX_assets", errors)
    boss_rows = unique_index(sheets.get("Bosses", []), "boss_id", "Bosses", errors)

    events_by_skill: dict[str, list[dict[str, Any]]] = defaultdict(list)
    compiled_events: dict[str, dict[str, Any]] = {}
    for event_id, row in event_rows.items():
        event = camel_mapping(row)
        event["id"] = event.pop("eventId")
        event["skillId"] = clean_string(event.get("skillId"))
        event["animationId"] = clean_string(event.pop("animId", None))
        event["refId"] = clean_string(event.get("refId"))
        event["loop"] = boolean(event.get("loop"))
        for field in ("frameStart", "frameEnd", "delayMs"):
            event[field] = number(event.get(field), f"Skill_events[{event_id}].{field}", errors, allow_none=True)
        compiled_events[event_id] = event
        events_by_skill[event["skillId"]].append(event)

    effects_by_skill: dict[str, list[dict[str, Any]]] = defaultdict(list)
    compiled_effects: dict[str, dict[str, Any]] = {}
    for effect_id, row in effect_rows.items():
        effect = camel_mapping(row)
        effect["id"] = effect.pop("effectId")
        for field in ("force", "amount", "durationMs", "maxStacks"):
            effect[field] = number(effect.get(field), f"Skill_effects[{effect_id}].{field}", errors, allow_none=True)
        compiled_effects[effect_id] = effect
        effects_by_skill[effect.get("skillId")].append(effect)

    compiled_projectiles: dict[str, dict[str, Any]] = {}
    for projectile_id, row in projectile_rows.items():
        projectile = camel_mapping(row, path_fields={"image_path"})
        projectile["id"] = projectile.pop("projectileId")
        for field in (
            "displayWidth",
            "displayHeight",
            "spawnForwardOffset",
            "spawnSideOffset",
            "speed",
            "colliderWidth",
            "colliderHeight",
            "lifetimeMs",
            "maxDistance",
            "pierceCount",
        ):
            projectile[field] = number(
                projectile.get(field),
                f"Projectiles[{projectile_id}].{field}",
                errors,
                allow_none=True,
            )
        for field in ("collideWorld", "destroyOnWorld", "destroyOnHit"):
            projectile[field] = boolean(projectile.get(field))
        compiled_projectiles[projectile_id] = projectile

    frames_by_animation: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_index, row in enumerate(sheets.get("Animation_frames", []), start=2):
        animation_id = clean_string(row.get("anim_id"))
        raw_frame = row.get("frame")
        raw_path = row.get("image_path")
        raw_duration = row.get("duration_ms")
        if not animation_id:
            errors.append(f"Animation_frames!{row_index}: blank anim_id")
            continue
        if is_placeholder(raw_frame) or is_placeholder(raw_path) or is_placeholder(raw_duration):
            warnings.append(f"Animation {animation_id} has TBD frame data")
            continue
        frame_index = number(raw_frame, f"Animation_frames!{row_index}.frame", errors)
        duration_ms = number(raw_duration, f"Animation_frames!{row_index}.duration_ms", errors)
        if frame_index is None or duration_ms is None:
            continue
        visual_scale = number(
            row.get("visual_scale"),
            f"Animation_frames!{row_index}.visual_scale",
            errors,
            allow_none=True,
        )
        frames_by_animation[animation_id].append(
            {
                "index": int(frame_index),
                "durationMs": duration_ms,
                "visualScale": visual_scale if visual_scale is not None else 1,
                "assetId": f"{animation_id}:{int(frame_index)}",
                "imagePath": normalize_path(raw_path),
                "notes": clean_string(row.get("notes")),
            }
        )

    direction_modes = {2: "horizontal_mirror", 4: "cardinal_4"}
    compiled_animations: dict[str, dict[str, Any]] = {}
    for animation_id, row in animation_rows.items():
        frames = sorted(frames_by_animation.get(animation_id, []), key=lambda item: item["index"])
        if frames:
            expected = list(range(frames[-1]["index"] + 1))
            actual = [frame["index"] for frame in frames]
            if actual != expected:
                errors.append(f"Animations[{animation_id}]: non-contiguous frames {actual}")
        else:
            warnings.append(f"Animation {animation_id} has no usable frames")
        mode_value = row.get("direction_mode")
        mode_number = int(mode_value) if isinstance(mode_value, (int, float)) else mode_value
        animation = {
            "id": animation_id,
            "ownerId": clean_string(row.get("owner_id")),
            "action": clean_string(row.get("action")),
            "loop": boolean(row.get("loop")),
            "loopStartFrame": number(
                row.get("loop_start_frame"),
                f"Animations[{animation_id}].loop_start_frame",
                errors,
                allow_none=True,
            ),
            "loopEndFrame": number(
                row.get("loop_end_frame"),
                f"Animations[{animation_id}].loop_end_frame",
                errors,
                allow_none=True,
            ),
            "directionMode": direction_modes.get(mode_number, str(mode_number)),
            "displayScale": number(
                row.get("display_scale"),
                f"Animations[{animation_id}].display_scale",
                errors,
                allow_none=True,
            )
            or 1,
            "required": boolean(row.get("required")),
            "frames": frames,
            "totalDurationMs": sum(frame["durationMs"] for frame in frames),
            "notes": clean_string(row.get("notes")),
        }
        compiled_animations[animation_id] = animation

    compiled_skills: dict[str, dict[str, Any]] = {}
    for skill_id, row in skill_rows.items():
        params_ref = clean_ref(row.get("params"))
        params = workbook_sheet_parameters(sheets, params_ref, f"Skills[{skill_id}]", errors)
        events = sorted(events_by_skill.get(skill_id, []), key=lambda item: item["id"])
        effects = sorted(effects_by_skill.get(skill_id, []), key=lambda item: item["id"])
        if clean_string(params.get("shape")) == "arc":
            for field in ("originForwardOffset", "originOffsetX", "originOffsetY", "angleDeg", "radius"):
                params[field] = number(
                    params.get(field),
                    f"{params_ref}[{skill_id}].{field}",
                    errors,
                )
            origin_reference = clean_string(params.get("origin")) or "actor"
            target_reference = clean_string(params.get("targetReference")) or "actor"
            if origin_reference not in ACTOR_REFERENCE_POINTS:
                errors.append(
                    f"{params_ref}[{skill_id}].origin: unsupported actor reference {origin_reference!r}"
                )
            if target_reference not in ACTOR_REFERENCE_POINTS:
                errors.append(
                    f"{params_ref}[{skill_id}].targetReference: unsupported actor reference "
                    f"{target_reference!r}"
                )
            params["origin"] = origin_reference
            params["targetReference"] = target_reference
        skill = {
            "id": skill_id,
            "name": clean_string(row.get("name_cn")),
            "target": clean_string(row.get("target")),
            "animationId": clean_string(row.get("anim_id")),
            "handler": infer_skill_handler(skill_id, params, events),
            "hitFrame": number(
                row.get("hit_frame"),
                f"Skills[{skill_id}].hit_frame",
                errors,
                allow_none=True,
            ),
            "castRange": {
                "min": number(row.get("cast_range_min"), f"Skills[{skill_id}].cast_range_min", errors),
                "max": number(row.get("cast_range_max"), f"Skills[{skill_id}].cast_range_max", errors),
            },
            "cooldownMs": number(row.get("cooldown_ms"), f"Skills[{skill_id}].cooldown_ms", errors),
            "weight": number(row.get("weight"), f"Skills[{skill_id}].weight", errors),
            "damage": number(row.get("damage"), f"Skills[{skill_id}].damage", errors),
            "params": params,
            "eventIds": [event["id"] for event in events],
            "effectIds": [effect["id"] for effect in effects],
            "chain": clean_string(row.get("chain")),
            "superArmor": boolean(row.get("super_armor")),
            "notes": clean_string(row.get("notes")),
        }
        if skill["castRange"]["max"] < skill["castRange"]["min"]:
            errors.append(f"Skills[{skill_id}]: cast_range_max must be >= cast_range_min")
        compiled_skills[skill_id] = skill

    compiled_vfx: dict[str, dict[str, Any]] = {}
    for vfx_id, row in vfx_rows.items():
        vfx = camel_mapping(row, path_fields={"image_path"})
        vfx["id"] = vfx.pop("vfxId")
        for field in ("durationMs", "scale", "alpha", "lineWidth", "radius"):
            vfx[field] = number(vfx.get(field), f"VFX_assets[{vfx_id}].{field}", errors, allow_none=True)
        vfx["loop"] = boolean(vfx.get("loop"))
        vfx["required"] = boolean(vfx.get("required"))
        compiled_vfx[vfx_id] = vfx

    compiled_sfx: dict[str, dict[str, Any]] = {}
    for sfx_id, row in sfx_rows.items():
        sfx = camel_mapping(row, path_fields={"path"})
        sfx["id"] = sfx.pop("sfxId")
        sfx["loop"] = boolean(sfx.get("loop"))
        sfx["required"] = boolean(sfx.get("required"))
        sfx["volume"] = number(sfx.get("volume"), f"SFX_assets[{sfx_id}].volume", errors, allow_none=True)
        sfx["maxInstances"] = number(
            sfx.get("maxInstances"),
            f"SFX_assets[{sfx_id}].max_instances",
            errors,
            allow_none=True,
        )
        if is_placeholder(sfx.get("path")):
            warnings.append(f"SFX {sfx_id} has no file path")
            sfx["path"] = None
        compiled_sfx[sfx_id] = sfx

    phase_rows_by_boss: dict[str, list[dict[str, Any]]] = defaultdict(list)
    phase_skill_rows: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for row_index, row in enumerate(sheets.get("Phase_skills", []), start=2):
        boss_id = clean_string(row.get("boss_id"))
        phase_value = number(row.get("phase"), f"Phase_skills!{row_index}.phase", errors)
        if not boss_id or phase_value is None:
            continue
        phase_skill_rows[(boss_id, int(phase_value))].append(row)

    for row_index, row in enumerate(sheets.get("Boss_phases", []), start=2):
        boss_id = clean_string(row.get("boss_id"))
        if not boss_id:
            errors.append(f"Boss_phases!{row_index}: blank boss_id")
            continue
        phase_value = number(row.get("phase"), f"Boss_phases!{row_index}.phase", errors)
        if phase_value is None:
            continue
        phase = int(phase_value)
        skills = []
        for skill_row in phase_skill_rows.get((boss_id, phase), []):
            if not boolean(skill_row.get("enabled")):
                continue
            skills.append(
                {
                    "skillId": clean_string(skill_row.get("skill_id")),
                    "weight": number(
                        skill_row.get("weight_override"),
                        f"Phase_skills[{boss_id}:{phase}].weight_override",
                        errors,
                        allow_none=True,
                    ),
                    "notes": clean_string(skill_row.get("notes")),
                }
            )
        phase_rows_by_boss[boss_id].append(
            {
                "phase": phase,
                "hpMinPct": number(row.get("hp_min_pct"), f"Boss_phases[{boss_id}:{phase}].hp_min_pct", errors),
                "hpMaxPct": number(row.get("hp_max_pct"), f"Boss_phases[{boss_id}:{phase}].hp_max_pct", errors),
                "minInclusive": boolean(row.get("min_inclusive")),
                "maxInclusive": boolean(row.get("max_inclusive")),
                "windupScale": number(
                    row.get("windup_scale"),
                    f"Boss_phases[{boss_id}:{phase}].windup_scale",
                    errors,
                    allow_none=True,
                )
                or 1,
                "transitionAnimationId": clean_string(row.get("transition_anim_id")),
                "transitionVfxId": clean_string(row.get("transition_vfx_id")),
                "transitionSfxId": clean_string(row.get("transition_sfx_id")),
                "skills": skills,
                "notes": clean_string(row.get("notes")),
            }
        )

    compiled_bosses: dict[str, dict[str, Any]] = {}
    for boss_id, row in boss_rows.items():
        phases = sorted(phase_rows_by_boss.get(boss_id, []), key=lambda item: item["phase"])
        boss = {
            "id": boss_id,
            "name": clean_string(row.get("name_cn")),
            "maxHp": number(row.get("max_hp"), f"Bosses[{boss_id}].max_hp", errors),
            "moveSpeed": number(row.get("move_speed"), f"Bosses[{boss_id}].move_speed", errors),
            "aggroRange": number(row.get("aggro_range"), f"Bosses[{boss_id}].aggro_range", errors),
            "display": {
                "width": number(row.get("display_width"), f"Bosses[{boss_id}].display_width", errors),
                "height": number(row.get("display_height"), f"Bosses[{boss_id}].display_height", errors),
                "anchorX": number(row.get("sprite_anchor_x"), f"Bosses[{boss_id}].sprite_anchor_x", errors),
                "anchorY": number(row.get("sprite_anchor_y"), f"Bosses[{boss_id}].sprite_anchor_y", errors),
            },
            "visualBounds": {
                "top": number(row.get("sprite_anchor_y"), f"Bosses[{boss_id}].sprite_anchor_y", errors),
            },
            "body": {
                "shape": clean_string(row.get("body_shape")),
                "width": number(row.get("body_width"), f"Bosses[{boss_id}].body_width", errors),
                "height": number(row.get("body_height"), f"Bosses[{boss_id}].body_height", errors),
                "offsetX": number(row.get("body_offset_x"), f"Bosses[{boss_id}].body_offset_x", errors),
                "offsetY": number(row.get("body_offset_y"), f"Bosses[{boss_id}].body_offset_y", errors),
            },
            "hurtbox": {
                "shape": clean_string(row.get("hurtbox_shape")),
                "width": number(row.get("hurtbox_width"), f"Bosses[{boss_id}].hurtbox_width", errors),
                "height": number(row.get("hurtbox_height"), f"Bosses[{boss_id}].hurtbox_height", errors),
                "offsetX": number(row.get("hurtbox_offset_x"), f"Bosses[{boss_id}].hurtbox_offset_x", errors),
                "offsetY": number(row.get("hurtbox_offset_y"), f"Bosses[{boss_id}].hurtbox_offset_y", errors),
            },
            "initialPhase": int(number(row.get("initial_phase"), f"Bosses[{boss_id}].initial_phase", errors) or 1),
            "animations": {
                "idle": clean_string(row.get("idle_anim_id")),
                "move": clean_string(row.get("move_anim_id")),
                "hurt": clean_string(row.get("hurt_anim_id")),
                "death": clean_string(row.get("death_anim_id")),
                "phaseTransition": clean_string(row.get("phase_transition_anim_id")),
            },
            "phases": phases,
            "notes": clean_string(row.get("notes")),
        }
        if not phases:
            errors.append(f"Bosses[{boss_id}]: no phases")
        compiled_bosses[boss_id] = boss

    validate_references(
        compiled_skills,
        compiled_events,
        compiled_effects,
        compiled_projectiles,
        compiled_animations,
        compiled_vfx,
        compiled_sfx,
        compiled_bosses,
        errors,
        warnings,
    )
    validate_asset_paths(
        compiled_projectiles,
        compiled_animations,
        compiled_vfx,
        warnings,
        errors,
    )

    return {
        "_meta": {
            "schemaVersion": CONFIG_SCHEMA_VERSION,
            "generator": "tools/build_boss_config.py",
            "generatorVersion": GENERATOR_VERSION,
            "source": source_path.name,
            "sourceSha256": sha256_file(source_path),
            "warnings": sorted(set(warnings)),
        },
        "bosses": compiled_bosses,
        "skills": compiled_skills,
        "skillEvents": compiled_events,
        "skillEffects": compiled_effects,
        "projectiles": compiled_projectiles,
        "animations": compiled_animations,
        "vfx": compiled_vfx,
        "sfx": compiled_sfx,
    }


def validate_references(
    skills: dict[str, dict[str, Any]],
    events: dict[str, dict[str, Any]],
    effects: dict[str, dict[str, Any]],
    projectiles: dict[str, dict[str, Any]],
    animations: dict[str, dict[str, Any]],
    vfx: dict[str, dict[str, Any]],
    sfx: dict[str, dict[str, Any]],
    bosses: dict[str, dict[str, Any]],
    errors: list[str],
    warnings: list[str],
) -> None:
    for skill_id, skill in skills.items():
        if skill["animationId"] not in animations:
            errors.append(f"Skills[{skill_id}]: unknown animation {skill['animationId']}")
        if skill["chain"] and skill["chain"] not in skills:
            errors.append(f"Skills[{skill_id}]: unknown chain skill {skill['chain']}")
        if skill["handler"] == "unsupported":
            warnings.append(f"Skill {skill_id} has unsupported handler")

    for event_id, event in events.items():
        if event["skillId"] not in skills:
            errors.append(f"Skill_events[{event_id}]: unknown skill {event['skillId']}")
        if event["animationId"] and event["animationId"] not in animations:
            errors.append(f"Skill_events[{event_id}]: unknown animation {event['animationId']}")
        ref_id = event.get("refId")
        event_type = event.get("eventType")
        registry: dict[str, Any] | None = None
        if event_type == "vfx":
            registry = vfx
        elif event_type == "sfx":
            registry = sfx
        elif event_type == "projectile_spawn":
            registry = projectiles
        elif event_type == "camera_shake":
            registry = vfx
        if registry is not None and ref_id not in registry:
            errors.append(f"Skill_events[{event_id}]: unknown {event_type} ref_id={ref_id}")

    for effect_id, effect in effects.items():
        if effect.get("skillId") not in skills:
            errors.append(f"Skill_effects[{effect_id}]: unknown skill {effect.get('skillId')}")

    for boss_id, boss in bosses.items():
        for role, animation_id in boss["animations"].items():
            if animation_id and animation_id not in animations:
                errors.append(f"Bosses[{boss_id}].animations.{role}: unknown {animation_id}")
        previous_max = None
        for phase in boss["phases"]:
            if phase["hpMinPct"] > phase["hpMaxPct"]:
                errors.append(f"Bosses[{boss_id}] phase {phase['phase']}: hp min exceeds max")
            if previous_max is not None and phase["hpMaxPct"] > previous_max:
                warnings.append(f"Bosses[{boss_id}] phase order is not descending by HP")
            previous_max = phase["hpMaxPct"]
            for phase_skill in phase["skills"]:
                if phase_skill["skillId"] not in skills:
                    errors.append(
                        f"Bosses[{boss_id}] phase {phase['phase']}: unknown skill {phase_skill['skillId']}"
                    )


def expand_pattern_exists(path_text: str) -> bool:
    pattern = path_text.replace("{DIR}", "*").replace("{FRAME}", "*")
    return any(PROJECT_ROOT.glob(pattern))


def validate_asset_paths(
    projectiles: dict[str, dict[str, Any]],
    animations: dict[str, dict[str, Any]],
    vfx: dict[str, dict[str, Any]],
    warnings: list[str],
    errors: list[str],
) -> None:
    for projectile_id, projectile in projectiles.items():
        path_text = projectile.get("imagePath")
        if path_text and not is_placeholder(path_text) and not (PROJECT_ROOT / path_text).exists():
            errors.append(f"Projectile {projectile_id}: missing image {path_text}")

    for animation_id, animation in animations.items():
        for frame in animation["frames"]:
            path_text = frame["imagePath"]
            if not path_text:
                continue
            exists = expand_pattern_exists(path_text) if "{" in path_text else (PROJECT_ROOT / path_text).exists()
            if not exists:
                message = f"Animation {animation_id}: missing image {path_text}"
                if animation["required"]:
                    errors.append(message)
                else:
                    warnings.append(message)

    for vfx_id, item in vfx.items():
        path_text = item.get("imagePath")
        if not path_text or is_placeholder(path_text):
            continue
        exists = expand_pattern_exists(path_text) if "{" in path_text else (PROJECT_ROOT / path_text).exists()
        if not exists:
            message = f"VFX {vfx_id}: missing image {path_text}"
            if item["required"]:
                errors.append(message)
            else:
                warnings.append(message)


def manifest_animation(animation: dict[str, Any]) -> dict[str, Any]:
    frames = animation["frames"]
    result: dict[str, Any] = {
        "action": animation["action"],
        "directionMode": animation["directionMode"],
        "displayScale": animation["displayScale"],
        "required": animation["required"],
        "frameCount": len(frames),
    }
    has_direction_pattern = any("{DIR}" in (frame.get("imagePath") or "") for frame in frames)
    if not has_direction_pattern:
        result["frames"] = [
            {
                "assetId": frame["assetId"],
                "index": frame["index"],
                "path": frame["imagePath"],
            }
            for frame in frames
        ]
        return result

    source_directions: dict[str, list[dict[str, Any]]] = {}
    for direction in ("D", "R", "U", "L", "UL", "UR", "DL", "DR"):
        expanded = []
        for frame in frames:
            path_text = frame["imagePath"].replace("{DIR}", direction)
            if (PROJECT_ROOT / path_text).exists():
                expanded.append(
                    {
                        "assetId": f"{frame['assetId']}:{direction}",
                        "index": frame["index"],
                        "path": path_text,
                    }
                )
        if len(expanded) == len(frames) and expanded:
            source_directions[direction] = expanded

    directions = {
        direction: {"source": direction, "mirror": False}
        for direction in source_directions
    }
    if "R" in source_directions and "L" not in source_directions:
        directions["L"] = {"source": "R", "mirror": True}
    if "UR" in source_directions and "UL" not in source_directions:
        directions["UL"] = {"source": "UR", "mirror": True}
    if "DR" in source_directions and "DL" not in source_directions:
        directions["DL"] = {"source": "DR", "mirror": True}
    result["directions"] = directions
    result["framesByDirection"] = source_directions
    return result


def manifest_vfx(vfx: dict[str, Any]) -> dict[str, Any]:
    path_text = vfx.get("imagePath")
    result = {
        "kind": vfx["kind"],
        "path": path_text,
        "required": vfx["required"],
    }
    if not path_text or "{FRAME}" not in path_text:
        return result

    pattern = path_text.replace("{FRAME}", "*")
    prefix, suffix = path_text.split("{FRAME}", 1)
    discovered = []
    for path in sorted(PROJECT_ROOT.glob(pattern)):
        relative = path.relative_to(PROJECT_ROOT).as_posix()
        if not relative.startswith(prefix) or not relative.endswith(suffix):
            continue
        raw_index = relative[len(prefix) : len(relative) - len(suffix) if suffix else None]
        if raw_index.isdigit():
            discovered.append({"index": int(raw_index), "path": relative})
    result["frames"] = discovered
    return result


def compile_manifest(
    manifest_source: Path,
    config: dict[str, Any],
    errors: list[str],
) -> dict[str, Any]:
    if not manifest_source.exists():
        raise BuildFailure(f"Manifest source does not exist: {manifest_source}")
    try:
        source = json.loads(manifest_source.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise BuildFailure(f"Invalid source manifest: {exc}") from exc

    manifest = copy.deepcopy(source)
    manifest["version"] = 2
    manifest["_generated"] = {
        "generator": "tools/build_boss_config.py",
        "generatorVersion": GENERATOR_VERSION,
        "bossConfigSha256": config["_meta"]["sourceSha256"],
    }
    manifest["data"] = {"bossConfig": "assets/data/boss_config.json"}
    manifest["bosses"] = {}

    animations = config["animations"]
    for boss_id, boss in config["bosses"].items():
        boss_animations = {
            animation_id: manifest_animation(animation)
            for animation_id, animation in animations.items()
            if animation["ownerId"] == boss_id
        }
        manifest["bosses"][boss_id] = {
            "displaySize": boss["display"],
            "atlas": {
                "metadata": f"assets/atlas/{boss_id}.json",
                "required": False,
            },
            "animations": boss_animations,
        }

    manifest["projectiles"] = {
        projectile_id: {
            "path": projectile["imagePath"],
            "displaySize": [projectile["displayWidth"], projectile["displayHeight"]],
            "required": True,
        }
        for projectile_id, projectile in config["projectiles"].items()
    }
    manifest["vfx"] = {
        vfx_id: manifest_vfx(vfx)
        for vfx_id, vfx in config["vfx"].items()
    }
    manifest["sfx"] = {
        sfx_id: {
            "path": sfx.get("path"),
            "required": sfx["required"],
        }
        for sfx_id, sfx in config["sfx"].items()
    }
    if errors:
        raise BuildFailure("Cannot compile manifest while validation errors exist")
    return manifest


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def runtime_config(config: dict[str, Any]) -> dict[str, Any]:
    result = copy.deepcopy(config)
    for animation in result["animations"].values():
        for frame in animation["frames"]:
            frame.pop("imagePath", None)
    for projectile in result["projectiles"].values():
        projectile.pop("imagePath", None)
    for vfx in result["vfx"].values():
        vfx.pop("imagePath", None)
    for sfx in result["sfx"].values():
        sfx.pop("path", None)
    return result


def atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(content)
    temporary.replace(path)


def csv_bytes(headers: list[str], rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header)
            if isinstance(value, bool):
                value = "TRUE" if value else "FALSE"
            elif value is None:
                value = ""
            values.append(value)
        writer.writerow(values)
    return buffer.getvalue().encode("utf-8-sig")


def write_csv_mirrors(
    workbook: Any,
    sheets: dict[str, list[dict[str, Any]]],
    output_dir: Path,
) -> None:
    for worksheet in workbook.worksheets:
        rows = sheets.get(worksheet.title, [])
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [clean_string(value) or "" for value in first_row]
        headers = [header for header in headers if header]
        safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", worksheet.title).strip("_") or "sheet"
        atomic_write(output_dir / f"{safe_name}.csv", csv_bytes(headers, rows))


def main() -> int:
    args = parse_args()
    errors: list[str] = []
    warnings: list[str] = []
    try:
        workbook, sheets = read_workbook(args.source.resolve())
        validate_columns(sheets, errors)
        config = compile_config(sheets, args.source.resolve(), errors, warnings)
        manifest = compile_manifest(args.manifest_source.resolve(), config, errors)
    except BuildFailure as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    if errors:
        print("Configuration validation failed:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 2

    for warning in sorted(set(warnings)):
        print(f"WARNING: {warning}")

    if args.check:
        print(
            f"OK: {len(config['bosses'])} boss(es), {len(config['skills'])} skill(s), "
            f"{len(config['animations'])} animation(s), {len(warnings)} warning(s)"
        )
        return 0

    atomic_write(args.config_out.resolve(), json_bytes(runtime_config(config)))
    atomic_write(args.manifest_out.resolve(), json_bytes(manifest))
    if not args.no_csv:
        write_csv_mirrors(workbook, sheets, args.csv_dir.resolve())
    print(f"Wrote {args.config_out.resolve()}")
    print(f"Wrote {args.manifest_out.resolve()}")
    if not args.no_csv:
        print(f"Wrote CSV mirrors to {args.csv_dir.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
